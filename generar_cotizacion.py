#!/usr/bin/env python3
"""Genera la cotización ASINMET en formato Word (.docx)"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ── Estilos base ──
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(10.5)

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)  # Azul oscuro

# ── Helpers ──
def set_cell_shading(cell, color):
    """Aplica color de fondo a una celda."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def add_table(doc, headers, rows, col_widths=None):
    """Crea tabla con encabezado coloreado."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Encabezados
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, '1B3A5C')

    # Filas
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
            if r_idx % 2 == 1:
                set_cell_shading(cell, 'EDF2F7')

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    return table

def add_kv_table(doc, pairs):
    """Tabla clave-valor de 2 columnas."""
    table = doc.add_table(rows=len(pairs), cols=2)
    table.style = 'Table Grid'
    for i, (k, v) in enumerate(pairs):
        ck = table.rows[i].cells[0]
        cv = table.rows[i].cells[1]
        ck.text = k
        cv.text = v
        for run in ck.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9.5)
        for run in cv.paragraphs[0].runs:
            run.font.size = Pt(9.5)
        set_cell_shading(ck, 'EDF2F7')
        ck.width = Cm(5)
        cv.width = Cm(12)
    return table


# ════════════════════════════════════════════════
# PORTADA
# ════════════════════════════════════════════════
doc.add_paragraph('')
doc.add_paragraph('')
doc.add_paragraph('')

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('COTIZACIÓN DE SERVICIOS')
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Implementación Odoo 19 Enterprise')
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x4A, 0x6F, 0xA5)

doc.add_paragraph('')

client = doc.add_paragraph()
client.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = client.add_run('ASINMET')
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)

client2 = doc.add_paragraph()
client2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = client2.add_run('Asociación de Industriales Metalúrgicos\nde la Provincia de Mendoza')
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph('')
doc.add_paragraph('')

add_kv_table(doc, [
    ('Proveedor', 'YAGUVEN C.G.'),
    ('CUIT Cliente', '30-55085607-3'),
    ('Domicilio', 'Lateral Norte Acceso Este 868, Guaymallén, Mendoza'),
    ('Fecha', '24/02/2026'),
    ('Validez', '30 días corridos'),
    ('Referencia', 'Reuniones del 13/02/2026 y 24/02/2026'),
])

doc.add_page_break()

# ════════════════════════════════════════════════
# 1. OBJETIVO
# ════════════════════════════════════════════════
doc.add_heading('1. OBJETIVO', level=1)

doc.add_paragraph(
    'Implementar Odoo 19 Enterprise como plataforma única de gestión para ASINMET, '
    'reemplazando los sistemas actuales por una solución integrada, accesible desde '
    'cualquier lugar y con eliminación de procesos en papel.'
)

add_table(doc,
    ['Sistema actual', 'Reemplazado por'],
    [
        ['Sistema SOL (2013)', 'Módulo de Medicina Laboral a medida'],
        ['VIXIS / Bixis (2011)', 'Contabilidad y Facturación electrónica AFIP'],
        ['Google Calendar', 'Agenda integrada con turnos y recursos'],
        ['Google Forms', 'Encuestas de satisfacción nativas'],
        ['Planillas manuales', 'CRM, cobranzas, pagos, reclamos integrados'],
    ],
    col_widths=[7, 10]
)

doc.add_paragraph('')

# ════════════════════════════════════════════════
# 2. ALCANCE
# ════════════════════════════════════════════════
doc.add_heading('2. ALCANCE', level=1)

# Etapa 1
doc.add_heading('ETAPA 1 — Core + Módulo Médico (Prioridad Alta)', level=2)

add_table(doc,
    ['#', 'Entregable', 'Detalle'],
    [
        ['1.1', 'Setup Odoo.sh', 'Instalación, configuración compañía, logo, datos fiscales, cuentas bancarias'],
        ['1.2', 'Contabilidad Argentina', 'Plan de cuentas (305 cuentas), configuración Exenta, conexión AFIP para Factura C electrónica con QR y CAE'],
        ['1.3', 'Módulo asinmet_medical', 'Desarrollo a medida: 15 modelos, 7 estados de workflow, agenda inteligente, ejecución de estudios, DDJJ (SRT 37/2010), certificado PDF con firma electrónica, emails automáticos, crons, 4 grupos de seguridad'],
        ['1.4', 'CRM', 'Pipeline comercial según PE-CO 01, seguimiento de oportunidades, KPIs'],
        ['1.5', 'Contactos', 'Carga de empresas socias (200+), clientes (700+), pacientes y profesionales'],
        ['1.6', 'Facturación', 'Facturas C electrónicas AFIP, envío por email, cobros, conciliación bancaria'],
        ['1.7', 'Agenda de turnos', 'Sedes, recursos (laboratorio, consultorios), calendarios, slots por capacidad'],
        ['1.8', 'Migración inicial', 'Saldos de apertura contable + base de empresas y contactos activos'],
    ],
    col_widths=[1, 4, 12]
)

doc.add_paragraph('')

# Etapa 2
doc.add_heading('ETAPA 2 — Expansión (Prioridad Media)', level=2)

add_table(doc,
    ['#', 'Entregable', 'Detalle'],
    [
        ['2.1', 'Comisiones comerciales', 'Reglas por vendedor, servicio, cliente nuevo'],
        ['2.2', 'Inventario de insumos', 'Stock de consumibles médicos (gel, jeringas, EPP), consumo por servicio'],
        ['2.3', 'Firma electrónica mejorada', 'Odoo Sign para profesionales externos vía portal'],
        ['2.4', 'Portal empresas', 'Descarga de certificados, consulta de facturas y pagos'],
    ],
    col_widths=[1, 5, 11]
)

doc.add_paragraph('')

# Etapa 3
doc.add_heading('ETAPA 3 — Complementarios (Prioridad Baja)', level=2)

add_table(doc,
    ['#', 'Entregable', 'Detalle'],
    [
        ['3.1', 'Eventos', 'Cursos del Centro Tecnológico (soldadura, capacitaciones), inscripciones online'],
        ['3.2', 'Encuestas', 'Satisfacción de pacientes (ML-F-01) y empresas (CO-F-07)'],
        ['3.3', 'Email Marketing', 'Campañas de difusión de servicios'],
        ['3.4', 'Sitio Web', 'Presencia web, publicación de servicios y cursos'],
        ['3.5', 'Activos fijos', 'Inventario de equipamiento del Centro Tecnológico'],
    ],
    col_widths=[1, 5, 11]
)

doc.add_page_break()

# ════════════════════════════════════════════════
# 3. MÓDULO MEDICINA LABORAL
# ════════════════════════════════════════════════
doc.add_heading('3. MÓDULO MEDICINA LABORAL — Detalle', level=1)

doc.add_paragraph(
    'El módulo asinmet_medical es un desarrollo 100% a medida que cubre el circuito completo '
    'del examen preocupacional según normativa SRT vigente (Res. 37/2010, Res. 53/2025).'
)

doc.add_heading('Flujo operativo', level=3)

flow = doc.add_paragraph()
flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = flow.add_run(
    'Solicitud → Orden de Examen → Asignación de turnos → Check-in\n'
    '→ Estudios en paralelo (laboratorio, ECG, Rx, audiometría, psicología...)\n'
    '→ Completitud automática → Auditoría médica → Dictamen\n'
    '→ Certificado PDF con firma → Entrega digital → Facturación'
)
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)

doc.add_paragraph('')

doc.add_heading('Componentes del módulo', level=3)

add_table(doc,
    ['Componente', 'Cantidad', 'Detalle'],
    [
        ['Modelos de datos', '15', 'Orden, estudios, resultados, certificados, agenda, DDJJ, recursos'],
        ['Templates de examen', '6', 'Preocupacional estándar/completo/con esfuerzo, periódico, egreso, ausencia'],
        ['Especialidades médicas', '10', 'Cardiología, audiología, radiología, neurología, psicología, etc.'],
        ['Reportes PDF', '2', 'Certificado de aptitud + Orden de control interno'],
        ['Emails automáticos', '5', 'Confirmación, reasignación, recordatorio, sin cupo, etc.'],
        ['Tareas programadas', '2', 'Recordatorio de turnos + vencimiento de reservas'],
        ['Grupos de seguridad', '4', 'Recepción → Técnico → Auditor → Gerente (jerárquicos)'],
        ['Dictámenes SRT', '4', 'Apto / Apto afecciones leves / Apto restricciones / No apto'],
    ],
    col_widths=[5, 2, 10]
)

doc.add_paragraph('')

doc.add_heading('Firma electrónica (Etapa 1)', level=3)
doc.add_paragraph(
    'Imagen holográfica del profesional almacenada en su perfil + timestamp + hash del documento. '
    'Validez legal como firma electrónica simple (Ley 25.506, Art. 5). '
    'Equivalente a la práctica actual de escaneo de firma.'
)

doc.add_page_break()

# ════════════════════════════════════════════════
# 4. USUARIOS
# ════════════════════════════════════════════════
doc.add_heading('4. USUARIOS Y LICENCIAS', level=1)

add_table(doc,
    ['Rol', 'Tipo', 'Licencia'],
    [
        ['Gerencia (admin)', 'Interno', 'Incluido en suscripción'],
        ['Jefa de Administración', 'Interno', 'Incluido en suscripción'],
        ['Contable', 'Interno', 'Incluido en suscripción'],
        ['Comercial', 'Interno', 'Incluido en suscripción'],
        ['Encargada Medicina', 'Interno', 'Incluido en suscripción'],
        ['Profesionales médicos (20-30)', 'Portal', 'Sin costo'],
        ['Empresas clientes (700+)', 'Portal', 'Sin costo'],
    ],
    col_widths=[6, 3, 5]
)

p = doc.add_paragraph('')
run = p.add_run('Total: 5 usuarios internos pagos + portal ilimitado sin costo adicional.')
run.bold = True

# ════════════════════════════════════════════════
# 5. CRONOGRAMA
# ════════════════════════════════════════════════
doc.add_heading('5. CRONOGRAMA', level=1)

add_table(doc,
    ['Etapa', 'Período', 'Hito clave'],
    [
        ['Etapa 1', 'Semana 1-4', 'Setup + Contabilidad + Facturación operativa (mes 1)'],
        ['', 'Semana 5-8', 'Módulo médico en producción'],
        ['', 'Semana 9-10', 'CRM + Agenda + Migración + Capacitación'],
        ['Etapa 2', 'Semana 11-16', 'Comisiones + Inventario + Firma + Portal'],
        ['Etapa 3', 'Semana 17-22', 'Eventos + Encuestas + Marketing + Web'],
    ],
    col_widths=[3, 4, 10]
)

p = doc.add_paragraph('')
run = p.add_run(
    'Compromiso: Contabilidad y facturación electrónica AFIP operativas dentro del primer mes, '
    'generando valor productivo inmediato.'
)
run.bold = True
run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)

doc.add_page_break()

# ════════════════════════════════════════════════
# 6. INVERSIÓN
# ════════════════════════════════════════════════
doc.add_heading('6. INVERSIÓN', level=1)

# A. Odoo
doc.add_heading('A. Suscripción Odoo.sh (pago anual a Odoo S.A.)', level=2)

add_table(doc,
    ['Concepto', 'Costo'],
    [
        ['Plataforma Odoo.sh (1 worker + staging + dev)', 'USD 72/mes'],
        ['5 usuarios Enterprise', 'USD 31,20/usuario/mes (*)'],
        ['Subtotal mensual', 'USD 228/mes'],
        ['Subtotal anual', 'USD 2.736/año'],
        ['Descuento 20% primer año (partner YAGUVEN)', '- USD 547'],
        ['TOTAL PRIMER AÑO', 'USD 2.189'],
    ],
    col_widths=[10, 5]
)

p = doc.add_paragraph(
    '(*) Precios estimados sujetos a cotización oficial de Odoo S.A. '
    'Incluyen hosting, SSL, backups automáticos, dominio personalizado, '
    'soporte Odoo y todas las actualizaciones.'
)
p.runs[0].font.size = Pt(9)
p.runs[0].font.italic = True

doc.add_paragraph('')

# B. Implementación
doc.add_heading('B. Implementación YAGUVEN', level=2)

add_table(doc,
    ['Concepto', 'Incluye', 'Inversión'],
    [
        ['Etapa 1 — Core + Módulo Médico',
         'Setup Odoo.sh, contabilidad AR, facturación AFIP, '
         'desarrollo módulo asinmet_medical (15 modelos), CRM, '
         'agenda de turnos, migración de datos, capacitación (5 sesiones), '
         'documentación, soporte post go-live 1 mes',
         'USD ________'],
        ['Etapa 2 — Expansión',
         'Comisiones, inventario insumos, firma electrónica Odoo Sign, '
         'portal empresas y profesionales, capacitación adicional',
         'USD ________'],
        ['Etapa 3 — Complementarios',
         'Eventos centro tecnológico, encuestas, email marketing, '
         'sitio web, activos fijos',
         'USD ________'],
    ],
    col_widths=[4, 9, 4]
)

p = doc.add_paragraph('')
run = p.add_run('Forma de pago: 3 cuotas mensuales iguales al inicio de cada mes por etapa contratada.')
run.bold = True

doc.add_paragraph('')

# C. Soporte
doc.add_heading('C. Soporte Post-Implementación (opcional)', level=2)

add_table(doc,
    ['Concepto', 'Detalle', 'Inversión'],
    [
        ['Soporte mensual (8 hs/mes)',
         'Resolución de incidencias, ajustes menores, consultas funcionales, actualizaciones de seguridad',
         'USD ________/mes'],
        ['Hora adicional',
         'Fuera del pack mensual, con presupuesto previo',
         'USD ________/hora'],
    ],
    col_widths=[5, 8, 4]
)

doc.add_page_break()

# ════════════════════════════════════════════════
# 7. CAPACITACIÓN
# ════════════════════════════════════════════════
doc.add_heading('7. CAPACITACIÓN', level=1)

add_table(doc,
    ['Sesión', 'Audiencia', 'Contenido', 'Duración'],
    [
        ['1. Introducción general', 'Todos los usuarios', 'Navegación Odoo, contactos, chatter, portal', '1,5 hs'],
        ['2. Contabilidad y facturación', 'Administración + Contable', 'Facturas AFIP, cobros, conciliación bancaria, reportes', '2 hs'],
        ['3. Módulo médico — Recepción', 'Recepcionista + Encargada', 'Órdenes de examen, turnos, check-in, DDJJ', '2 hs'],
        ['4. Módulo médico — Clínica', 'Técnicos + Auditor', 'Ejecución de estudios, carga de resultados, certificados', '2 hs'],
        ['5. CRM y gestión comercial', 'Comercial + Gerencia', 'Pipeline, oportunidades, actividades, KPIs', '1,5 hs'],
    ],
    col_widths=[5, 4, 5, 2]
)

p = doc.add_paragraph('')
run = p.add_run('Total: 9 horas de capacitación incluidas en Etapa 1.')
run.bold = True

doc.add_paragraph(
    'Modalidad: Virtual (Google Meet / Zoom) o presencial en sede ASINMET.\n'
    'Material: Manuales de usuario, parametrizador y desarrollador incluidos.'
)

# ════════════════════════════════════════════════
# 8. ACOMPAÑAMIENTO
# ════════════════════════════════════════════════
doc.add_heading('8. ACOMPAÑAMIENTO POST GO-LIVE', level=1)

add_table(doc,
    ['Período', 'Incluido en', 'Actividades'],
    [
        ['Mes 1 post go-live', 'Etapa 1',
         'Soporte prioritario, ajustes funcionales, resolución de incidencias, '
         'acompañamiento diario los primeros 5 días hábiles'],
        ['Mes 2 en adelante', 'Contrato de soporte',
         'Soporte por ticket/email/WhatsApp, SLA 24 hs hábiles, ajustes menores, consultas'],
    ],
    col_widths=[4, 4, 9]
)

p = doc.add_paragraph('')
p.add_run(
    'Escalamiento: ').bold = True
p.add_run(
    'Requerimientos nuevos (fuera del alcance original) se cotizan por separado con presupuesto previo.'
)

doc.add_page_break()

# ════════════════════════════════════════════════
# 9. CONDICIONES
# ════════════════════════════════════════════════
doc.add_heading('9. CONDICIONES GENERALES', level=1)

conditions = [
    ('Propiedad del código',
     'El módulo asinmet_medical es propiedad de ASINMET. '
     'YAGUVEN retiene derecho de reutilización del framework '
     '(no de los datos ni reglas de negocio específicas).'),
    ('Entorno',
     'Odoo.sh gestionado por YAGUVEN como partner técnico. '
     'ASINMET es titular de la cuenta y la suscripción.'),
    ('Responsabilidades de ASINMET',
     '• Designar un referente por área para relevamiento y validación\n'
     '• Proveer datos para migración en formato acordado\n'
     '• Disponibilidad para sesiones de capacitación\n'
     '• Testeo y validación del módulo médico antes del go-live\n'
     '• Listado de equipos médicos (marca, modelo, versión) para evaluar integración'),
    ('Responsabilidades de YAGUVEN',
     '• Desarrollo, configuración y despliegue del sistema\n'
     '• Capacitación y documentación\n'
     '• Soporte post go-live según lo contratado\n'
     '• Backups y mantenimiento del entorno Odoo.sh'),
    ('Exclusiones',
     '• Integración directa con equipos médicos (se evalúa caso por caso en Etapa 2)\n'
     '• Migración de datos históricos completos del Sistema SOL (solo saldos y contactos activos)\n'
     '• Desarrollo de app mobile nativa\n'
     '• Firma digital certificada (requiere proveedor externo, se evalúa en Etapa 2-3)'),
    ('Vigencia',
     '30 días corridos desde la fecha de emisión.'),
]

for title, text in conditions:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)
    p.add_run(text)

# ════════════════════════════════════════════════
# 10. PRÓXIMOS PASOS
# ════════════════════════════════════════════════
doc.add_heading('10. PRÓXIMOS PASOS', level=1)

add_table(doc,
    ['#', 'Acción', 'Responsable', 'Plazo'],
    [
        ['1', 'Revisión de esta cotización con equipo directivo', 'ASINMET', '1 semana'],
        ['2', 'Solicitud de cotización formal Odoo.sh a Odoo S.A.', 'YAGUVEN', 'Inmediato'],
        ['3', 'Aprobación y firma', 'ASINMET + YAGUVEN', 'A convenir'],
        ['4', 'Pago primera cuota + inicio Etapa 1', 'ASINMET', 'Según firma'],
        ['5', 'Setup Odoo.sh + contabilidad operativa', 'YAGUVEN', 'Mes 1'],
        ['6', 'Equipo médico prueba prototipo y envía feedback', 'ASINMET', 'Paralelo'],
    ],
    col_widths=[1, 8, 4, 3]
)

doc.add_paragraph('')
doc.add_paragraph('')

# ── Firma ──
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('YAGUVEN C.G.')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.add_run('Horacio — Dirección técnica | Ricardo — Dirección funcional')

doc.add_paragraph('')

note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = note.add_run(
    'Cotización elaborada en base al relevamiento de las reuniones del 13/02/2026 y 24/02/2026, '
    'procedimientos ISO 9001 de ASINMET (PE-AD 01, PE-CO 01, PE-ML 01), '
    'normativa SRT vigente y prototipo funcional demostrado en entorno de testing.'
)
run.font.size = Pt(8.5)
run.font.italic = True
run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# ── Guardar ──
output_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    'COTIZACION_ASINMET_YAGUVEN_2026.docx'
)
doc.save(output_path)
print(f'Documento generado: {output_path}')
